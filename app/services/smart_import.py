"""Importador único: detecta o tipo de arquivo e delega para o parser certo.

Formatos reconhecidos
- Nubank conta: OFX ou CSV (Data, Valor, Identificador, Descrição)
- Nubank cartão: OFX ou CSV da fatura (date, title, amount[, category])
- B3 Área do Investidor (custódia NuInvest e outras corretoras), XLSX:
  Posição (ações, FIIs, ETFs, BDRs, renda fixa, Tesouro), Negociação e Movimentação
- Modelos CSV do app: posições, cotações, operações e renda fixa/caixinhas
"""
from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from app.db import transaction
from app.services.accounts import link_accounts
from app.services.categories import INTERNAL_CATEGORIES, categorize
from app.services.formatting import money_to_cents, quantity_to_micros
from app.services.importers import (
    _as_date, _csv_rows, _digest, _finish_run, _instrument, _start_run,
    import_ofx, import_positions_csv, import_quotes_csv, normalize_ticker,
)

NUINVEST_ALIAS = "Nubank / NuInvest"
FILENAME_DATE_RE = re.compile(r"(\d{4})[-_.]?(\d{2})[-_.]?(\d{2})")


@dataclass
class ImportResult:
    kind: str
    count: int
    message: str


def _norm(text: object) -> str:
    decomposed = unicodedata.normalize("NFKD", str(text or "").strip().casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _date_from_filename(filename: str) -> str | None:
    match = FILENAME_DATE_RE.search(filename)
    if not match:
        return None
    try:
        return date(*(int(part) for part in match.groups())).isoformat()
    except ValueError:
        return None


def _money(value: object) -> int | None:
    if value is None or (isinstance(value, str) and value.strip() in {"", "-", "—"}):
        return None
    try:
        return money_to_cents(value)
    except ValueError:
        return None


def _qty(value: object) -> int | None:
    if value is None or (isinstance(value, str) and value.strip() in {"", "-", "—"}):
        return None
    try:
        return quantity_to_micros(value)
    except ValueError:
        return None


def broker_account_name(institution: object) -> str:
    """Custódia da NuInvest usa o mesmo apelido do Open Finance para consolidar sem duplicar."""
    text = _norm(institution)
    if not text or any(key in text for key in ("nu invest", "nuinvest", "easynvest", "nubank", "nu pagamentos")):
        return NUINVEST_ALIAS
    return str(institution).strip().title()[:60]


def _account(connection, provider: str, external_key: str, name: str, account_type: str, institution: str) -> int:
    connection.execute(
        "INSERT INTO financial_account(institution, account_name, account_type, currency, provider, external_key) "
        "VALUES (?, ?, ?, 'BRL', ?, ?) "
        "ON CONFLICT(provider, external_key) DO UPDATE SET account_name = excluded.account_name",
        (institution, name, account_type, provider, external_key),
    )
    return int(connection.execute(
        "SELECT id FROM financial_account WHERE provider = ? AND external_key = ?", (provider, external_key)
    ).fetchone()[0])


def _single_ofx_account(connection, account_type: str) -> tuple[int, str] | None:
    found = connection.execute(
        "SELECT id, external_key FROM financial_account WHERE provider = 'ofx' AND account_type = ?",
        (account_type,),
    ).fetchall()
    return (int(found[0][0]), str(found[0][1])) if len(found) == 1 else None


# ---------------------------------------------------------------- entrada

def import_file(filename: str, content: bytes) -> ImportResult:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".ofx", ".qfx", ".ofc"}:
        count = import_ofx(content)
        return ImportResult("Extrato OFX Nubank", count, f"{count} movimentações")
    if suffix in {".xlsx", ".xlsm"}:
        return _import_workbook(filename, content)
    if suffix in {".csv", ".txt"}:
        return _import_csv(filename, content)
    if content[:200].lstrip().upper().startswith((b"OFXHEADER", b"<OFX", b"<?XML")):
        count = import_ofx(content)
        return ImportResult("Extrato OFX Nubank", count, f"{count} movimentações")
    raise ValueError("Formato não reconhecido. Envie OFX, CSV ou XLSX.")


def _import_csv(filename: str, content: bytes) -> ImportResult:
    records = _csv_rows(content)
    if not records:
        raise ValueError("O CSV não contém linhas de dados.")
    columns = {_norm(column) for column in records[0]}
    if {"data", "valor", "identificador"} <= columns:
        return _import_nuconta_csv(records)
    if {"date", "title", "amount"} <= columns:
        return _import_card_csv(filename, records)
    if {"ticker", "quantity", "as_of_date"} <= columns:
        count = import_positions_csv(content)
        return ImportResult("Posições (CSV)", count, f"{count} posições")
    if {"ticker", "trade_date", "close"} <= columns:
        count = import_quotes_csv(content)
        return ImportResult("Cotações (CSV)", count, f"{count} fechamentos")
    if {"date", "ticker", "type", "quantity"} <= columns:
        return _import_operations_csv(records)
    if {"name", "gross_value"} <= columns:
        return _import_fixed_income_csv(records)
    raise ValueError(
        "CSV não reconhecido. Use o extrato/fatura do Nubank ou um dos modelos disponíveis na página de importação."
    )


# ---------------------------------------------------------------- Nubank conta (CSV)

def _import_nuconta_csv(records: list[dict[str, str]]) -> ImportResult:
    rows = [{_norm(key): value for key, value in row.items()} for row in records]
    with transaction() as connection:
        run_id = _start_run(connection, "CSV Nubank conta")
        existing = _single_ofx_account(connection, "BANK")
        if existing:
            account_id, external_key = existing
        else:
            external_key = _digest("nubank-conta-csv")
            account_id = _account(connection, "ofx", external_key, "Nubank Conta", "BANK", "Nubank")
        # O Identificador do CSV é o mesmo FITID do OFX: reimportar em qualquer formato não duplica.
        last_date = None
        for row in rows:
            tx_date = _as_date(row["data"])
            amount = money_to_cents(row["valor"])
            description = (row.get("descricao") or "Movimentação Nubank").strip()[:250]
            identifier = (row.get("identificador") or "").strip() or _digest(f"{tx_date}|{description}|{amount}")
            connection.execute(
                "INSERT INTO cash_transaction(account_id, sync_run_id, source, external_id, transaction_date, "
                "description, amount_cents, currency, category) VALUES (?, ?, 'ofx', ?, ?, ?, ?, 'BRL', ?) "
                "ON CONFLICT(source, external_id) DO UPDATE SET sync_run_id = excluded.sync_run_id",
                (account_id, run_id, f"{external_key}:{identifier}", tx_date, description, amount, categorize(description)),
            )
            last_date = max(last_date or tx_date, tx_date)
        link_accounts(connection)
        derived = _derive_balance(connection, account_id, run_id, last_date)
        _finish_run(connection, run_id, len(rows), f"{len(rows)} movimentação(ões) da conta Nubank.")
    note = "" if derived else " · saldo não incluso no CSV: importe um OFX uma vez para ancorar o saldo"
    return ImportResult("Extrato Nubank (CSV)", len(rows), f"{len(rows)} movimentações{note}")


def _derive_balance(connection, account_id: int, run_id: int, until: str | None) -> bool:
    """Com um saldo-âncora de OFX/Open Finance, projeta o saldo somando as movimentações posteriores."""
    if not until:
        return False
    anchor = connection.execute(
        "SELECT as_of_date, balance_cents FROM account_balance_snapshot "
        "WHERE account_id = ? AND source <> 'derived' ORDER BY as_of_date DESC LIMIT 1",
        (account_id,),
    ).fetchone()
    if not anchor or anchor["as_of_date"] >= until:
        return bool(anchor)
    movement = connection.execute(
        "SELECT COALESCE(SUM(amount_cents), 0) FROM cash_transaction "
        "WHERE account_id = ? AND status = 'POSTED' AND transaction_date > ? AND transaction_date <= ?",
        (account_id, anchor["as_of_date"], until),
    ).fetchone()[0]
    connection.execute(
        "INSERT INTO account_balance_snapshot(account_id, sync_run_id, source, as_of_date, balance_cents) "
        "VALUES (?, ?, 'derived', ?, ?) ON CONFLICT(account_id, source, as_of_date) DO UPDATE SET "
        "balance_cents = excluded.balance_cents, sync_run_id = excluded.sync_run_id, imported_at = CURRENT_TIMESTAMP",
        (account_id, run_id, until, int(anchor["balance_cents"]) + int(movement)),
    )
    return True


# ---------------------------------------------------------------- Nubank cartão (CSV da fatura)

def _import_card_csv(filename: str, records: list[dict[str, str]]) -> ImportResult:
    rows = [{_norm(key): value for key, value in row.items()} for row in records]
    dates = [_as_date(row["date"]) for row in rows]
    statement_ref = _date_from_filename(filename) or max(dates)
    with transaction() as connection:
        run_id = _start_run(connection, "CSV fatura Nubank")
        existing = _single_ofx_account(connection, "CREDIT")
        if existing:
            account_id, external_key = existing
        else:
            external_key = _digest("nubank-cartao-csv")
            account_id = _account(connection, "ofx", external_key, "Nubank Cartão", "CREDIT", "Nubank")
        seen: Counter[str] = Counter()
        bill_total = 0
        for row, tx_date in zip(rows, dates):
            title = (row.get("title") or "Compra no cartão").strip()[:250]
            # Na fatura, compra é positiva; aqui toda saída é negativa, como na conta.
            amount = -money_to_cents(row["amount"])
            category = categorize(title, row.get("category"))
            if category not in INTERNAL_CATEGORIES:
                bill_total += amount
            fingerprint = f"{tx_date}|{title}|{amount}"
            seen[fingerprint] += 1
            external_id = f"{external_key}:{statement_ref}:{_digest(fingerprint + f'|{seen[fingerprint]}')}"
            connection.execute(
                "INSERT INTO cash_transaction(account_id, sync_run_id, source, external_id, transaction_date, "
                "description, amount_cents, currency, category, statement_ref) "
                "VALUES (?, ?, 'nubank_csv', ?, ?, ?, ?, 'BRL', ?, ?) "
                "ON CONFLICT(source, external_id) DO UPDATE SET sync_run_id = excluded.sync_run_id, "
                "category = excluded.category",
                (account_id, run_id, external_id, tx_date, title, amount, category, statement_ref),
            )
        connection.execute(
            "INSERT INTO account_balance_snapshot(account_id, sync_run_id, source, as_of_date, balance_cents) "
            "VALUES (?, ?, 'nubank_csv', ?, ?) ON CONFLICT(account_id, source, as_of_date) DO UPDATE SET "
            "balance_cents = excluded.balance_cents, sync_run_id = excluded.sync_run_id, imported_at = CURRENT_TIMESTAMP",
            (account_id, run_id, statement_ref, min(bill_total, 0)),
        )
        link_accounts(connection)
        _finish_run(connection, run_id, len(rows), f"Fatura {statement_ref}: {len(rows)} lançamento(s).")
    return ImportResult("Fatura Nubank (CSV)", len(rows), f"fatura {statement_ref} · {len(rows)} lançamentos")


# ---------------------------------------------------------------- Operações e renda fixa (modelos CSV)

OPERATION_TYPES = {
    "buy": "BUY", "compra": "BUY", "c": "BUY",
    "sell": "SELL", "venda": "SELL", "v": "SELL",
    "dividend": "DIVIDEND", "dividendo": "DIVIDEND",
    "jcp": "JCP", "juros sobre capital proprio": "JCP",
    "rendimento": "INCOME", "income": "INCOME",
    "split": "SPLIT", "desdobro": "SPLIT", "bonus": "BONUS", "bonificacao": "BONUS",
}


def _import_operations_csv(records: list[dict[str, str]]) -> ImportResult:
    rows = [{_norm(key): value for key, value in row.items()} for row in records]
    with transaction() as connection:
        run_id = _start_run(connection, "CSV operações")
        seen: Counter[str] = Counter()
        for row in rows:
            ticker = normalize_ticker(row["ticker"])
            if not ticker:
                raise ValueError(f"Ticker B3 inválido: {row['ticker']!r}.")
            event_type = OPERATION_TYPES.get(_norm(row["type"]))
            if not event_type:
                raise ValueError(f"Tipo de operação desconhecido: {row['type']!r}. Use BUY, SELL, DIVIDEND, JCP, RENDIMENTO, SPLIT ou BONUS.")
            event_date = _as_date(row["date"])
            quantity = abs(_qty(row.get("quantity")) or 0)
            price = _money(row.get("price"))
            fees = abs(_money(row.get("fees")) or 0)
            amount = _money(row.get("amount"))
            if amount is None and price is not None:
                amount = int(round(quantity * price / 1_000_000))
            if amount is not None:
                amount = abs(amount) + (fees if event_type == "BUY" else -fees if event_type == "SELL" else 0)
            if event_type in {"DIVIDEND", "JCP", "INCOME"} and amount is None:
                raise ValueError(f"{ticker}: informe amount (ou price) para proventos.")
            signed_qty = -quantity if event_type == "SELL" else quantity if event_type in {"BUY", "SPLIT", "BONUS"} else 0
            account_name = (row.get("account_name") or NUINVEST_ALIAS).strip()
            account_id = _account(connection, "manual_ops", _digest("ops:" + account_name.casefold()), account_name, "INVESTMENT", "Importação manual")
            instrument_id = _instrument(connection, ticker)
            fingerprint = f"{account_name}|{event_date}|{ticker}|{event_type}|{signed_qty}|{amount}"
            seen[fingerprint] += 1
            connection.execute(
                "INSERT INTO investment_event(account_id, instrument_id, sync_run_id, source, external_id, event_date, "
                "event_type, quantity_micros, amount_cents, unit_price_cents, description) "
                "VALUES (?, ?, ?, 'manual_ops', ?, ?, ?, ?, ?, ?, 'Operação importada por CSV') "
                "ON CONFLICT(source, external_id) DO UPDATE SET sync_run_id = excluded.sync_run_id",
                (account_id, instrument_id, run_id, _digest(fingerprint + f"|{seen[fingerprint]}"),
                 event_date, event_type, signed_qty, amount, price),
            )
        _finish_run(connection, run_id, len(rows), f"{len(rows)} operação(ões) importada(s).")
    return ImportResult("Operações (CSV)", len(rows), f"{len(rows)} operações")


def _product_type(name: str, hint: str = "") -> str:
    text = _norm(hint or name)
    for label, keys in (
        ("Caixinha", ("caixinha", "rdb")), ("CDB", ("cdb",)), ("LCI", ("lci",)), ("LCA", ("lca",)),
        ("Tesouro", ("tesouro",)), ("CRI/CRA", ("cri", "cra")), ("Debênture", ("debent",)),
        ("Fundo", ("fundo", "fic", "fim")), ("Previdência", ("previd", "pgbl", "vgbl")), ("LC", ("lc ",)),
    ):
        if any(key in text for key in keys):
            return label
    return hint.strip().title() or "Renda fixa"


def _write_fixed_income(connection, account_id: int, run_id: int, source: str, product_key: str, as_of: str, data: dict[str, Any]) -> None:
    connection.execute(
        "INSERT INTO fixed_income_snapshot(account_id, sync_run_id, source, product_key, product_type, name, issuer, "
        "indexer, rate, maturity_date, as_of_date, invested_cents, gross_value_cents, net_value_cents) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?) "
        "ON CONFLICT(product_key, source, as_of_date) DO UPDATE SET account_id = excluded.account_id, "
        "sync_run_id = excluded.sync_run_id, product_type = excluded.product_type, name = excluded.name, "
        "issuer = excluded.issuer, indexer = excluded.indexer, rate = excluded.rate, maturity_date = excluded.maturity_date, "
        "invested_cents = excluded.invested_cents, gross_value_cents = excluded.gross_value_cents, "
        "net_value_cents = excluded.net_value_cents, imported_at = CURRENT_TIMESTAMP",
        (account_id, run_id, source, product_key, data["product_type"], data["name"][:120], data.get("issuer"),
         data.get("indexer"), data.get("rate"), data.get("maturity_date"), as_of, data.get("invested_cents"),
         data["gross_value_cents"], data.get("net_value_cents")),
    )


def _import_fixed_income_csv(records: list[dict[str, str]]) -> ImportResult:
    rows = [{_norm(key): value for key, value in row.items()} for row in records]
    with transaction() as connection:
        run_id = _start_run(connection, "CSV renda fixa")
        for row in rows:
            name = row["name"].strip()
            if not name:
                raise ValueError("Informe o nome de cada produto de renda fixa.")
            gross = _money(row["gross_value"])
            if gross is None:
                raise ValueError(f"{name}: gross_value inválido.")
            account_name = (row.get("account_name") or NUINVEST_ALIAS).strip()
            account_id = _account(connection, "manual", _digest("fixed-income:" + account_name.casefold()), account_name, "INVESTMENT", "Importação manual")
            maturity = row.get("maturity") or row.get("maturity_date")
            _write_fixed_income(connection, account_id, run_id, "manual", _digest(f"rf:{account_name.casefold()}:{name.casefold()}"),
                _as_date(row.get("as_of_date") or date.today().isoformat()), {
                    "product_type": _product_type(name, row.get("type", "")), "name": name,
                    "issuer": row.get("issuer") or None, "indexer": row.get("indexer") or None,
                    "rate": row.get("rate") or None, "maturity_date": _as_date(maturity) if maturity else None,
                    "invested_cents": _money(row.get("invested")), "gross_value_cents": gross,
                    "net_value_cents": _money(row.get("net_value")),
                })
        _finish_run(connection, run_id, len(rows), f"{len(rows)} produto(s) de renda fixa.")
    return ImportResult("Renda fixa (CSV)", len(rows), f"{len(rows)} produtos")


# ---------------------------------------------------------------- B3 Área do Investidor (XLSX)

EQUITY_SHEETS = {"acoes": "Ação", "bdr": "BDR", "etf": "ETF", "fundo de investimento": "FII", "fii": "FII"}


def _sheet_records(sheet) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows[:8]) if row and sum(1 for cell in row if isinstance(cell, str) and cell.strip()) >= 3),
        None,
    )
    if header_index is None:
        return []
    header = [_norm(cell) for cell in rows[header_index]]
    output = []
    for row in rows[header_index + 1:]:
        if not row or not any(cell not in (None, "") for cell in row):
            continue
        record = {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}
        first = _norm(next(iter(record.values()), ""))
        if first.startswith("total"):
            continue
        output.append(record)
    return output


def _cell_date(value: Any) -> str:
    if hasattr(value, "date") and callable(value.date):
        return value.date().isoformat()
    return _as_date(value)


def _import_workbook(filename: str, content: bytes) -> ImportResult:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ImportError as exc:
        raise ValueError("Instale as dependências (openpyxl) executando run.bat novamente.") from exc
    except Exception as exc:
        raise ValueError("Não foi possível abrir a planilha XLSX.") from exc
    sheets = {_norm(name): _sheet_records(workbook[name]) for name in workbook.sheetnames}
    workbook.close()
    all_columns = {column for records in sheets.values() for record in records[:1] for column in record}
    if {"data do negocio", "tipo de movimentacao", "codigo de negociacao"} <= all_columns:
        return _import_b3_trades(sheets)
    if {"entrada/saida", "movimentacao", "produto"} <= all_columns:
        return _import_b3_movements(sheets)
    if any(name in EQUITY_SHEETS or name in {"renda fixa", "tesouro direto"} for name in sheets):
        return _import_b3_position(filename, sheets)
    raise ValueError("Planilha não reconhecida. Exporte Posição, Negociação ou Movimentação na Área do Investidor da B3.")


def _b3_account(connection, institution: object) -> int:
    name = broker_account_name(institution)
    return _account(connection, "b3", _digest("b3:" + name.casefold()), name, "INVESTMENT", str(institution or "B3"))


def _import_b3_position(filename: str, sheets: dict[str, list[dict[str, Any]]]) -> ImportResult:
    as_of = _date_from_filename(filename) or date.today().isoformat()
    positions = fixed = 0
    with transaction() as connection:
        run_id = _start_run(connection, "B3 posição")
        seen_positions: set[tuple[int, int]] = set()
        seen_products: set[str] = set()
        for sheet_name, records in sheets.items():
            asset_class = EQUITY_SHEETS.get(sheet_name)
            for record in records:
                if asset_class:
                    ticker = normalize_ticker(record.get("codigo de negociacao"))
                    quantity = _qty(record.get("quantidade"))
                    if not ticker or quantity is None:
                        continue
                    account_id = _b3_account(connection, record.get("instituicao"))
                    product = str(record.get("produto") or "")
                    name = product.split(" - ", 1)[1].strip() if " - " in product else ticker
                    instrument_id = _instrument(connection, ticker, asset_class, "BRL", name[:80])
                    key = (account_id, instrument_id)
                    already = key in seen_positions
                    seen_positions.add(key)
                    connection.execute(
                        "INSERT INTO position_snapshot(account_id, instrument_id, sync_run_id, source, as_of_date, quantity_micros, reported_value_cents) "
                        "VALUES (?, ?, ?, 'b3', ?, ?, ?) ON CONFLICT(account_id, instrument_id, source, as_of_date) DO UPDATE SET "
                        + ("quantity_micros = position_snapshot.quantity_micros + excluded.quantity_micros, " if already else "quantity_micros = excluded.quantity_micros, ")
                        + "reported_value_cents = excluded.reported_value_cents, sync_run_id = excluded.sync_run_id, imported_at = CURRENT_TIMESTAMP",
                        (account_id, instrument_id, run_id, as_of, quantity, _money(record.get("valor atualizado"))),
                    )
                    close = _money(record.get("preco de fechamento"))
                    if close and close > 0:
                        connection.execute(
                            "INSERT INTO daily_quote(instrument_id, trade_date, close_cents, provider) VALUES (?, ?, ?, 'b3') "
                            "ON CONFLICT(instrument_id, trade_date, provider) DO UPDATE SET close_cents = excluded.close_cents, fetched_at = CURRENT_TIMESTAMP",
                            (instrument_id, as_of, close),
                        )
                    positions += 1
                elif sheet_name in {"renda fixa", "tesouro direto"}:
                    product = str(record.get("produto") or "").strip()
                    if not product:
                        continue
                    if sheet_name == "tesouro direto":
                        gross = _money(record.get("valor bruto")) or _money(record.get("valor atualizado"))
                        data = {"product_type": "Tesouro", "name": product, "issuer": "Tesouro Nacional",
                                "invested_cents": _money(record.get("valor aplicado")),
                                "net_value_cents": _money(record.get("valor liquido"))}
                    else:
                        gross = _money(record.get("valor atualizado curva")) or _money(record.get("valor atualizado mtm"))
                        data = {"product_type": _product_type(product), "name": product,
                                "issuer": str(record.get("emissor") or "") or None}
                    if gross is None:
                        continue
                    maturity = record.get("vencimento")
                    data.update({
                        "gross_value_cents": gross, "indexer": str(record.get("indexador") or "") or None,
                        "maturity_date": _cell_date(maturity) if maturity not in (None, "", "-") else None,
                    })
                    account_id = _b3_account(connection, record.get("instituicao"))
                    code = record.get("codigo") or record.get("codigo isin") or ""
                    product_key = _digest(f"b3rf:{account_id}:{code}:{product}:{data['maturity_date']}")
                    already = product_key in seen_products
                    seen_products.add(product_key)
                    if already:
                        connection.execute(
                            "UPDATE fixed_income_snapshot SET gross_value_cents = gross_value_cents + ? "
                            "WHERE product_key = ? AND source = 'b3' AND as_of_date = ?",
                            (gross, product_key, as_of),
                        )
                    else:
                        _write_fixed_income(connection, account_id, run_id, "b3", product_key, as_of, data)
                    fixed += 1
        # A posição da B3 é uma fotografia completa da custódia: o que sumiu foi vendido ou venceu.
        closed = _close_missing_b3(connection, run_id, as_of, seen_positions, seen_products)
        _finish_run(connection, run_id, positions + fixed, f"Posição B3 {as_of}: {positions} ativo(s), {fixed} título(s).")
    extra = f" · {closed} encerrado(s)" if closed else ""
    return ImportResult("Posição B3", positions + fixed, f"{positions} ativos listados · {fixed} títulos de renda fixa{extra} · data {as_of}")


def _close_missing_b3(connection, run_id: int, as_of: str, positions: set[tuple[int, int]], products: set[str]) -> int:
    closed = 0
    for row in connection.execute(
        "SELECT account_id, instrument_id, quantity_micros, as_of_date FROM v_latest_position_per_account "
        "WHERE source = 'b3' AND quantity_micros <> 0"
    ).fetchall():
        if (row["account_id"], row["instrument_id"]) in positions or row["as_of_date"] >= as_of:
            continue
        connection.execute(
            "INSERT OR REPLACE INTO position_snapshot(account_id, instrument_id, sync_run_id, source, as_of_date, quantity_micros) "
            "VALUES (?, ?, ?, 'b3', ?, 0)", (row["account_id"], row["instrument_id"], run_id, as_of),
        )
        closed += 1
    for row in connection.execute(
        "SELECT f.* FROM fixed_income_snapshot f WHERE f.source = 'b3' AND f.as_of_date = "
        "(SELECT MAX(as_of_date) FROM fixed_income_snapshot g WHERE g.product_key = f.product_key AND g.source = 'b3') "
        "AND f.gross_value_cents <> 0"
    ).fetchall():
        if row["product_key"] in products or row["as_of_date"] >= as_of:
            continue
        _write_fixed_income(connection, row["account_id"], run_id, "b3", row["product_key"], as_of, {
            "product_type": row["product_type"], "name": row["name"], "issuer": row["issuer"],
            "indexer": row["indexer"], "maturity_date": row["maturity_date"], "gross_value_cents": 0,
        })
        closed += 1
    return closed


def _import_b3_trades(sheets: dict[str, list[dict[str, Any]]]) -> ImportResult:
    count = 0
    with transaction() as connection:
        run_id = _start_run(connection, "B3 negociação")
        seen: Counter[str] = Counter()
        for records in sheets.values():
            for record in records:
                ticker = normalize_ticker(record.get("codigo de negociacao"))
                side = _norm(record.get("tipo de movimentacao"))
                quantity = _qty(record.get("quantidade"))
                if not ticker or quantity is None or side not in {"compra", "venda"}:
                    continue
                event_type = "BUY" if side == "compra" else "SELL"
                event_date = _cell_date(record.get("data do negocio"))
                amount = abs(_money(record.get("valor")) or 0)
                price = _money(record.get("preco"))
                account_id = _b3_account(connection, record.get("instituicao"))
                instrument_id = _instrument(connection, ticker)
                signed = -abs(quantity) if event_type == "SELL" else abs(quantity)
                fingerprint = f"{account_id}|{event_date}|{ticker}|{event_type}|{signed}|{amount}"
                seen[fingerprint] += 1
                connection.execute(
                    "INSERT INTO investment_event(account_id, instrument_id, sync_run_id, source, external_id, event_date, "
                    "event_type, quantity_micros, amount_cents, unit_price_cents, description) "
                    "VALUES (?, ?, ?, 'b3', ?, ?, ?, ?, ?, ?, 'Negociação B3') "
                    "ON CONFLICT(source, external_id) DO UPDATE SET sync_run_id = excluded.sync_run_id",
                    (account_id, instrument_id, run_id, _digest(fingerprint + f"|{seen[fingerprint]}"),
                     event_date, event_type, signed, amount, price),
                )
                count += 1
        _finish_run(connection, run_id, count, f"{count} negociação(ões) B3.")
    return ImportResult("Negociações B3", count, f"{count} negociações")


MOVEMENT_TYPES = {
    "dividendo": "DIVIDEND", "juros sobre capital proprio": "JCP", "rendimento": "INCOME",
    "desdobro": "SPLIT", "bonificacao em ativos": "BONUS", "grupamento": "SPLIT",
    "transferencia - liquidacao": "SETTLEMENT",
}


def _import_b3_movements(sheets: dict[str, list[dict[str, Any]]]) -> ImportResult:
    count = 0
    with transaction() as connection:
        run_id = _start_run(connection, "B3 movimentação")
        seen: Counter[str] = Counter()
        for records in sheets.values():
            for record in records:
                movement = MOVEMENT_TYPES.get(_norm(record.get("movimentacao")))
                product = str(record.get("produto") or "")
                ticker = normalize_ticker(product.split(" - ", 1)[0])
                if not movement or not ticker:
                    continue
                credit = _norm(record.get("entrada/saida")).startswith("credito")
                quantity = abs(_qty(record.get("quantidade")) or 0)
                amount = _money(record.get("valor da operacao"))
                price = _money(record.get("preco unitario"))
                if movement == "SETTLEMENT":
                    event_type, signed = ("BUY", quantity) if credit else ("SELL", -quantity)
                elif movement in {"SPLIT", "BONUS"}:
                    event_type, signed = movement, quantity if credit else -quantity
                else:
                    if not credit or not amount:
                        continue
                    event_type, signed = movement, 0
                event_date = _cell_date(record.get("data"))
                account_id = _b3_account(connection, record.get("instituicao"))
                instrument_id = _instrument(connection, ticker, "Ação", "BRL", product.split(" - ", 1)[-1].strip()[:80] or None)
                fingerprint = f"{account_id}|{event_date}|{ticker}|{event_type}|{signed}|{amount}"
                seen[fingerprint] += 1
                # Liquidações ficam em fonte própria: se houver o arquivo de Negociação, ele tem prioridade.
                source = "b3_mov" if movement == "SETTLEMENT" else "b3"
                connection.execute(
                    "INSERT INTO investment_event(account_id, instrument_id, sync_run_id, source, external_id, event_date, "
                    "event_type, quantity_micros, amount_cents, unit_price_cents, description) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?) "
                    "ON CONFLICT(source, external_id) DO UPDATE SET sync_run_id = excluded.sync_run_id",
                    (account_id, instrument_id, run_id, source, _digest("mov|" + fingerprint + f"|{seen[fingerprint]}"),
                     event_date, event_type, signed, abs(amount) if amount is not None else None, price,
                     str(record.get("movimentacao") or "")[:120]),
                )
                count += 1
        _finish_run(connection, run_id, count, f"{count} movimentação(ões) B3.")
    return ImportResult("Movimentação B3", count, f"{count} eventos (proventos, desdobros, liquidações)")

